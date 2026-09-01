"""HDFC savings/current account statement — the Excel export from HDFC's
own website ("Download Statement"), a different source and format from
the emailed PDF handled by hdfc_bank_account.py.

Unlike the emailed PDF (a genuine ruled table pdfplumber merges into one
row per page, requiring reconstruction), this is a clean tabular export
with one row per real transaction and no line-wrapping at all — columns
are Date, Narration, Chq./Ref.No., Value Dt, Withdrawal Amt.,
Deposit Amt., Closing Balance. Reading and interpreting the sheet are
kept as separate functions (read_workbook_rows does file I/O via openpyxl
or xlrd; _rows_to_transactions is pure and unit-testable with plain
lists) — the same read/normalize split used throughout this package.

Verified against four real statements spanning April 2021 - April 2025:
the two source libraries disagree on how they represent an empty
Withdrawal/Deposit cell (xlrd's .xls reader gives '', openpyxl's .xlsx
reader gives None) — both are handled here.
"""
from __future__ import annotations

import re
from pathlib import Path

HEADER_ROW = ["Date", "Narration", "Amount"]

_DATE_RE = re.compile(r"^\d{2}/\d{2}/\d{2,4}$")


def read_workbook_rows(path: Path) -> list[list]:
    """Reads every cell of the first sheet as a list of rows, values
    as-is (numbers stay numbers, not strings) — dispatches on file
    extension since HDFC's own site has been observed to export both the
    legacy binary .xls format and modern .xlsx depending on when the
    statement was downloaded.
    """
    suffix = Path(path).suffix.lower()
    if suffix == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        ws = wb[wb.sheetnames[0]]
        return [list(row) for row in ws.iter_rows(values_only=True)]
    if suffix == ".xls":
        import xlrd
        wb = xlrd.open_workbook(str(path))
        sheet = wb.sheet_by_index(0)
        return [sheet.row_values(i) for i in range(sheet.nrows)]
    raise ValueError(f"Unsupported spreadsheet format: {suffix!r} (expected .xls or .xlsx)")


def _find_header_row(rows: list[list]) -> int | None:
    """The real column header ("Date", "Narration", ...) sits after a
    variable-length preamble (account holder name/address/branch block) —
    found by content, not a fixed row index, since that preamble's length
    has been observed to vary between statements.
    """
    for i, row in enumerate(rows[:40]):
        text = " ".join(str(c) for c in row if c is not None)
        if "Narration" in text and "Date" in text:
            return i
    return None


def _rows_to_transactions(rows: list[list]) -> list[list[str]]:
    """Turns the raw sheet rows into (Date, Narration, signed Amount)
    rows, munim import's expected shape. Rows before/after the real
    transaction table (the preamble, the '****' mask row directly under
    the header, and the footer's legal/disclosure text) are skipped by
    requiring column 0 to look like a DD/MM/YY(YY) date — none of that
    surrounding text does.
    """
    header_idx = _find_header_row(rows)
    if header_idx is None:
        return []
    result = []
    for row in rows[header_idx + 1:]:
        if not row or row[0] is None:
            continue
        date_str = str(row[0]).strip()
        if not _DATE_RE.match(date_str):
            continue
        narration = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
        withdrawal = row[4] if len(row) > 4 else None
        deposit = row[5] if len(row) > 5 else None
        try:
            w = float(withdrawal) if withdrawal not in (None, "") else 0.0
            d = float(deposit) if deposit not in (None, "") else 0.0
        except (TypeError, ValueError):
            continue
        result.append([date_str, narration, f"{d - w:.2f}"])
    return result


def parse_hdfc_bank_excel(path: Path) -> list[list[str]]:
    """Reads and parses one HDFC bank-account Excel statement export in
    one step — the public entry point; read_workbook_rows and
    _rows_to_transactions exist separately so the parsing logic can be
    unit-tested without a real spreadsheet file on disk.
    """
    return _rows_to_transactions(read_workbook_rows(path))
